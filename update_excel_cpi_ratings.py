import openpyxl
from openpyxl.utils import get_column_letter

def update_cpi_ratings_in_excel():
    print("Loading EXCEL8EXAMPLE.xlsx...")
    wb = openpyxl.load_workbook('EXCEL8EXAMPLE.xlsx', data_only=False)
    ws = wb['CPI data']

    # Official 29-step scale from rating rule.docx (same structure as Interest Rate)
    cpi_scale = [
        (7.0, 3),
        (6.5, 4),
        (6.0, 5),
        (5.5, 6),
        (5.0, 10),
        (4.5, 9),
        (4.0, 8),
        (3.5, 7),
        (3.0, 6),
        (2.5, 5),
        (2.0, 4),
        (1.5, 3),
        (1.0, 2),
        (0.5, 1),
        (0.0, 0),
        (-0.5, -1),
        (-1.0, -2),
        (-1.5, -3),
        (-2.0, -4),
        (-2.5, -5),
        (-3.0, -6),
        (-3.5, -7),
        (-4.0, -8),
        (-4.5, -9),
        (-5.0, -10),
        (-5.5, -6),
        (-6.0, -5),
        (-6.5, -4),
        (-7.0, -3),
    ]

    print("Writing official 29-step CPI rating table to AY6:AZ34...")
    for idx, (val, rat) in enumerate(cpi_scale):
        row = 6 + idx
        ws[f'AY{row}'] = val
        ws[f'AZ{row}'] = rat

    # Find month columns in Row 43 (Header row for pairs & ratings)
    max_col = 2
    while ws.cell(43, max_col).value is not None:
        max_col += 1
    num_cols = max_col - 2
    print(f"Found {num_cols} month columns from Col B to {get_column_letter(max_col-1)}")

    # Rows 44 to 66 are the 23 FX pairs in CPI data
    # Each pair rating row r corresponds to differential row r - 29 (15 to 37)
    print("Updating rating formulas in Rows 44 to 66...")
    for r in range(44, 67):
        pair_name = ws.cell(r, 1).value
        r_diff = r - 29
        for c in range(2, max_col):
            col_letter = get_column_letter(c)
            formula = (
                f'=_xlfn.XLOOKUP('
                f'MIN(ABS($AY$6:$AY$34 - MEDIAN(MIN($AY$6:$AY$34), MAX($AY$6:$AY$34), {col_letter}{r_diff}))), '
                f'ABS($AY$6:$AY$34 - MEDIAN(MIN($AY$6:$AY$34), MAX($AY$6:$AY$34), {col_letter}{r_diff})), '
                f'$AZ$6:$AZ$34, "Error", 0)'
            )
            ws.cell(r, c).value = formula

    wb.save('EXCEL8EXAMPLE.xlsx')
    print("[SUCCESS] EXCEL8EXAMPLE.xlsx updated with official CPI rating rules from rating rule.docx!")

if __name__ == '__main__':
    update_cpi_ratings_in_excel()
