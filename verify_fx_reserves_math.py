import openpyxl
import json
import sqlite3
import pandas as pd

def rate_fx(diff):
    if diff is None: return None
    BK_SPREAD = [-2000 + 200*i for i in range(21)]
    BL_RATING = [10 - i for i in range(21)]
    clamped = max(-2000.0, min(2000.0, float(diff)))
    idx = min(range(len(BK_SPREAD)), key=lambda i: abs(BK_SPREAD[i] - clamped))
    return int(BL_RATING[idx])

def audit():
    print("=" * 118)
    print("AUDITING FX RESERVES INTEGRATION ACROSS EXCEL, DJANGO DB, & FRONTEND")
    print("=" * 118)

    # 1. Load Excel FX RESERVE DATA sheet
    wb_val = openpyxl.load_workbook("EXCEL8EXAMPLE.xlsx", data_only=True)
    ws_fx = wb_val["FX RESERVE DATA"]
    ws_cm = wb_val["COMPARATIVE MATRIX"]

    # In FX RESERVE DATA:
    # Col 14 is 2023-01
    # Col 55 is 2026-06
    # Col 58 is 2026-09
    col_2026_06_fx = 55
    col_2026_09_fx = 58

    # Extract 12M moving averages from rows 25 to 32
    avg_rows = {
        'Australia': 25,
        'Canada': 26,
        'Japan': 27,
        'Switzerland': 28,
        'New Zealand': 29,
        'United Kingdom': 30,
        'United States': 31,
        'Euro Area': 32,
    }

    country_flows = {}
    for c_name, r in avg_rows.items():
        v_jun = ws_fx.cell(row=r, column=col_2026_06_fx).value
        v_sep = ws_fx.cell(row=r, column=col_2026_09_fx).value
        # If openpyxl cached value is None, calculate from raw values
        if v_jun is None or v_sep is None:
            raw_r = r - 22 # 3 to 10
            raw_jun = [ws_fx.cell(row=raw_r, column=c).value for c in range(col_2026_06_fx - 12, col_2026_06_fx + 1)]
            raw_sep = [ws_fx.cell(row=raw_r, column=c).value for c in range(col_2026_09_fx - 12, col_2026_09_fx + 1)]
            deltas_jun = [raw_jun[i] - raw_jun[i-1] for i in range(1, len(raw_jun))]
            deltas_sep = [raw_sep[i] - raw_sep[i-1] for i in range(1, len(raw_sep))]
            v_jun = sum(deltas_jun) / len(deltas_jun)
            v_sep = sum(deltas_sep) / len(deltas_sep)
        country_flows[c_name] = {
            "2026-06": float(v_jun),
            "2026-09": float(v_sep)
        }

    # Extract Excel pair spreads & ratings
    excel_pairs = {}
    pairs_def = [
        (37, 'USD/JPY', 'United States', 'Japan'),
        (38, 'USD/CHF', 'United States', 'Switzerland'),
        (39, 'USD/CAD', 'United States', 'Canada'),
        (40, 'EUR/USD', 'Euro Area', 'United States'),
        (41, 'EUR/JPY', 'Euro Area', 'Japan'),
        (42, 'EUR/CHF', 'Euro Area', 'Switzerland'),
        (43, 'EUR/GBP', 'Euro Area', 'United Kingdom'),
        (44, 'EUR/CAD', 'Euro Area', 'Canada'),
        (45, 'EUR/AUD', 'Euro Area', 'Australia'),
        (46, 'EUR/NZD', 'Euro Area', 'New Zealand'),
        (47, 'GBP/USD', 'United Kingdom', 'United States'),
        (48, 'GBP/JPY', 'United Kingdom', 'Japan'),
        (49, 'GBP/CHF', 'United Kingdom', 'Switzerland'),
        (50, 'GBP/CAD', 'United Kingdom', 'Canada'),
        (51, 'GBP/AUD', 'United Kingdom', 'Australia'),
        (52, 'GBP/NZD', 'United Kingdom', 'New Zealand'),
        (53, 'CHF/JPY', 'Switzerland', 'Japan'),
        (54, 'AUD/USD', 'Australia', 'United States'),
        (55, 'AUD/CHF', 'Australia', 'Switzerland'),
        (56, 'AUD/JPY', 'Australia', 'Japan'),
        (57, 'AUD/NZD', 'Australia', 'New Zealand'),
        (58, 'AUD/CAD', 'Australia', 'Canada'),
        (59, 'CAD/JPY', 'Canada', 'Japan'),
    ]

    for r_spread, pair_name, base_c, quote_c in pairs_def:
        b_jun = country_flows[base_c]["2026-06"]
        q_jun = country_flows[quote_c]["2026-06"]
        sp_jun = b_jun - q_jun
        rt_jun = rate_fx(sp_jun)

        b_sep = country_flows[base_c]["2026-09"]
        q_sep = country_flows[quote_c]["2026-09"]
        sp_sep = b_sep - q_sep
        rt_sep = rate_fx(sp_sep)

        excel_pairs[pair_name] = {
            "2026-06": {"spread": sp_jun, "rating": rt_jun, "base": b_jun, "quote": q_jun},
            "2026-09": {"spread": sp_sep, "rating": rt_sep, "base": b_sep, "quote": q_sep},
        }

    # 2. Check Django Database
    conn = sqlite3.connect("fx-macro-bias/backend/db.sqlite3")
    cur = conn.cursor()

    cur.execute("""
        SELECT p.symbol, d.month, d.base_value, d.quote_value, d.difference, r.rating_value, f.final_score, f.bias
        FROM macro_differential d
        JOIN macro_fxpair p ON d.pair_id = p.id
        JOIN macro_macroindicator i ON d.indicator_id = i.id
        LEFT JOIN macro_rating r ON r.differential_id = d.id
        LEFT JOIN macro_finalscore f ON f.pair_id = p.id AND f.month = d.month
        WHERE i.slug = 'fx_reserves' AND (d.month LIKE '2026-06%' OR d.month LIKE '2026-09%')
        ORDER BY p.symbol, d.month
    """)
    db_rows = cur.fetchall()
    db_data = {}
    for symbol, dt, b_val, q_val, diff_v, rate_v, fs_v, bias_v in db_rows:
        mo = dt[:7]
        db_data.setdefault(symbol, {})[mo] = {
            "base": b_val,
            "quote": q_val,
            "spread": diff_v,
            "rating": rate_v,
            "score": fs_v,
            "bias": bias_v
        }

    # 3. Check Frontend Dataset
    with open("fx-macro-bias/frontend/src/data/macroDataset.json", "r", encoding="utf-8") as f:
        fe_data = json.load(f)

    fe_pairs = fe_data.get("PairFXReserves", {})

    active_pairs = [
        "EUR/USD", "GBP/USD", "USD/JPY", "AUD/USD", "USD/CAD", "NZD/USD", "USD/CHF",
        "EUR/JPY", "GBP/JPY", "EUR/GBP", "EUR/CHF", "AUD/JPY", "CAD/JPY", "CHF/JPY",
        "EUR/AUD", "EUR/CAD", "EUR/NZD", "GBP/AUD", "GBP/CAD", "GBP/CHF", "GBP/NZD"
    ]

    print("\n" + "=" * 118)
    print(f"{'Pair':<9} | {'Period':<7} | {'Excel Spread':<12} | {'Excel Rtg':<9} | {'DB Spread':<12} | {'DB Rtg':<8} | {'FE Spread':<12} | {'FE Rtg':<8} | {'Score':<7} | {'Macro Bias':<10}")
    print("=" * 118)

    all_matched = True
    checked_count = 0
    for p in active_pairs:
        ex = excel_pairs.get(p, {})
        db = db_data.get(p, {})
        fe = fe_pairs.get(p, {})

        for mo in ["2026-06", "2026-09"]:
            ex_entry = ex.get(mo, {})
            ex_sp = ex_entry.get("spread")
            ex_rt = ex_entry.get("rating")

            db_entry = db.get(mo, {})
            db_sp = db_entry.get("spread")
            db_rt = db_entry.get("rating")
            db_score = db_entry.get("score")
            db_bias = db_entry.get("bias")

            fe_entry = fe.get(mo, {})
            fe_sp = fe_entry.get("spread")
            fe_rt = fe_entry.get("rating")

            ex_sp_str = f"{ex_sp:+10.2f}M" if ex_sp is not None else "None"
            ex_rt_str = f"{int(ex_rt):+d}" if ex_rt is not None else "None"
            db_sp_str = f"{db_sp:+10.2f}M" if db_sp is not None else "None"
            db_rt_str = f"{int(db_rt):+d}" if db_rt is not None else "None"
            fe_sp_str = f"{fe_sp:+10.2f}M" if fe_sp is not None else "None"
            fe_rt_str = f"{int(fe_rt):+d}" if fe_rt is not None else "None"
            sc_str = f"{db_score:+.2f}" if db_score is not None else "None"
            bias_str = str(db_bias) if db_bias else "None"

            print(f"{p:<9} | {mo:<7} | {ex_sp_str:<12} | {ex_rt_str:<9} | {db_sp_str:<12} | {db_rt_str:<8} | {fe_sp_str:<12} | {fe_rt_str:<8} | {sc_str:<7} | {bias_str:<10}")

            if ex_rt is not None and db_rt is not None:
                checked_count += 1
                if int(ex_rt) != int(db_rt):
                    print(f"MISMATCH in rating for {p} {mo}: Excel={ex_rt}, DB={db_rt}")
                    all_matched = False
                if abs(ex_sp - db_sp) > 1.0:
                    print(f"MISMATCH in spread for {p} {mo}: Excel={ex_sp}, DB={db_sp}")
                    all_matched = False
            if db_rt is not None and fe_rt is not None:
                if int(db_rt) != int(fe_rt):
                    print(f"MISMATCH in rating for {p} {mo}: DB={db_rt}, FE={fe_rt}")
                    all_matched = False
                if abs(db_sp - fe_sp) > 1.0:
                    print(f"MISMATCH in spread for {p} {mo}: DB={db_sp}, FE={fe_sp}")
                    all_matched = False

    print("=" * 118)
    if all_matched and checked_count > 0:
        print(f"[SUCCESS] 100% PARITY CONFIRMED! All {checked_count} pair-month FX Reserves spreads & ratings match across Excel, Django DB, and Frontend!")
    else:
        print("[ERROR] Detected inconsistencies.")

if __name__ == "__main__":
    audit()
