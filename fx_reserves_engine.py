import openpyxl
import pandas as pd
import numpy as np
import urllib.request
import json
import datetime
from pathlib import Path

# 8 G10 economies + Sweden & Norway
COUNTRIES = [
    'United States', 'Euro Area', 'Japan', 'United Kingdom',
    'Australia', 'Canada', 'Switzerland', 'New Zealand', 'Sweden', 'Norway'
]

PAIRS = [
    ('EUR/USD', 'Euro Area', 'United States'),
    ('USD/JPY', 'United States', 'Japan'),
    ('GBP/USD', 'United Kingdom', 'United States'),
    ('AUD/USD', 'Australia', 'United States'),
    ('USD/CAD', 'United States', 'Canada'),
    ('USD/CHF', 'United States', 'Switzerland'),
    ('NZD/USD', 'New Zealand', 'United States'),
    ('EUR/JPY', 'Euro Area', 'Japan'),
    ('GBP/JPY', 'United Kingdom', 'Japan'),
    ('EUR/GBP', 'Euro Area', 'United Kingdom'),
    ('EUR/CHF', 'Euro Area', 'Switzerland'),
    ('AUD/JPY', 'Australia', 'Japan'),
    ('CAD/JPY', 'Canada', 'Japan'),
    ('CHF/JPY', 'Switzerland', 'Japan'),
    ('EUR/AUD', 'Euro Area', 'Australia'),
    ('EUR/CAD', 'Euro Area', 'Canada'),
    ('EUR/NZD', 'Euro Area', 'New Zealand'),
    ('GBP/AUD', 'United Kingdom', 'Australia'),
    ('GBP/CAD', 'United Kingdom', 'Canada'),
    ('GBP/CHF', 'United Kingdom', 'Switzerland'),
    ('GBP/NZD', 'United Kingdom', 'New Zealand'),
    ('AUD/CAD', 'Australia', 'Canada'),
    ('AUD/CHF', 'Australia', 'Switzerland'),
    ('AUD/NZD', 'Australia', 'New Zealand'),
]

# Rating lookup grid from EXCEL8EXAMPLE.xlsx (Cols BK & BL)
BK_SPREAD = [-2000 + 200*i for i in range(21)]
BL_RATING = [10 - i for i in range(21)]

def rate_fx_spread(spread):
    if spread is None or pd.isna(spread):
        return 0
    clamped = max(-2000.0, min(2000.0, float(spread)))
    idx = min(range(len(BK_SPREAD)), key=lambda i: abs(BK_SPREAD[i] - clamped))
    return int(BL_RATING[idx])

def build_fx_dataset():
    """
    Builds authoritative FX Reserves dataset:
    - Reads raw data from EXCEL8EXAMPLE.xlsx (Cols B to BI, 2022-01 to 2026-12)
    - Enriches with official live IMF SDMX API data preserving full floating-point precision (NO ceil/floor/round)
    - Executes the 3-step calculation per 'Calcul du différentiel — FX Reserves Exchange.docx'
    """
    print("--- 1. Reading raw FX Reserves from EXCEL8EXAMPLE.xlsx ---")
    excel_path = Path(__file__).resolve().parent / "EXCEL8EXAMPLE.xlsx"
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    ws = wb['FX RESERVE DATA']

    excel_name_map = {
        'Australia': 'Australia',
        'Canada': 'Canada',
        'Japan': 'Japan',
        'Switzerland': 'Switzerland',
        'New Zealand': 'New Zealand',
        'United Kingdom': 'United Kingdom',
        'United States': 'United States',
        'EURO AREA': 'Euro Area',
    }

    dates = [ws.cell(row=2, column=c).value for c in range(2, 62)]
    date_strs = [d.strftime('%Y-%m') if hasattr(d, 'strftime') else str(d)[:7] for d in dates]

    country_series = {}
    for r in range(3, 11):
        raw_name = ws.cell(row=r, column=1).value
        std_name = excel_name_map.get(raw_name, raw_name)
        vals = [ws.cell(row=r, column=c).value for c in range(2, 62)]
        country_series[std_name] = pd.Series(vals, index=date_strs, dtype=float)

    # Fetch live IMF SDMX API data for exact precision where available
    print("--- 2. Querying Live IMF SDMX API (unrounded float data) ---")
    imf_code_map = {
        'USA': 'United States',
        'EA20': 'Euro Area',
        'G163': 'Euro Area',
        'JPN': 'Japan',
        'GBR': 'United Kingdom',
        'AUS': 'Australia',
        'CAN': 'Canada',
        'CHE': 'Switzerland',
        'NZL': 'New Zealand',
        'SWE': 'Sweden',
        'NOR': 'Norway',
    }
    try:
        codes = '+'.join(['USA', 'EA20', 'JPN', 'GBR', 'AUS', 'CAN', 'CHE', 'NZL', 'SWE', 'NOR'])
        url = f'https://api.imf.org/external/sdmx/3.0/data/dataflow/IMF.STA/IL/13.0.1/{codes}.RXF11FX_REVS.USD.M?c%5BTIME_PERIOD%5D=ge:2022-01-01&attributes=all&detail=full'
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0', 'Accept': 'application/json'})
        with urllib.request.urlopen(req, timeout=15) as resp:
            data = json.loads(resp.read().decode('utf-8'))
        s0 = data['data']['structures'][0]
        c_list = [c['id'] for c in s0['dimensions']['series'][0]['values']]
        periods = [t['value'] for t in s0['dimensions']['observation'][0]['values']]
        series = data['data']['dataSets'][0]['series']

        for s_key, s_data in series.items():
            code = c_list[int(s_key.split(':')[0])]
            std_name = imf_code_map.get(code)
            if not std_name: continue
            if std_name not in country_series:
                country_series[std_name] = pd.Series(index=date_strs, dtype=float)
            obs = s_data.get('observations', {})
            for k, v in obs.items():
                p_str = periods[int(k)].replace('-M', '-')
                if p_str in country_series[std_name].index and v[0] is not None:
                    # Exact float: USD Millions without floor or ceil
                    country_series[std_name].loc[p_str] = float(v[0]) / 1e6
        print("  -> Live IMF SDMX API data successfully incorporated.")
    except Exception as e:
        print(f"  -> Note: IMF SDMX API query ({e}). Using Excel baseline.")

    df_raw = pd.DataFrame(country_series)
    # Forward fill any missing months towards the end of 2026
    df_raw = df_raw.ffill()

    print("--- 3. Executing 3-Step FX Reserves Math ---")
    # Étape 1: Variation mensuelle (Reserves(N) - Reserves(N-1))
    df_delta = df_raw.diff()

    # Étape 2: Moyenne glissante sur 12 mois
    df_avg12 = df_delta.rolling(window=12, min_periods=12).mean()

    # Étape 3: Différentiel entre les pays (Base - Quote) & Rating
    pair_results = {}
    for pair_name, base_c, quote_c in PAIRS:
        if base_c in df_avg12.columns and quote_c in df_avg12.columns:
            base_flows = df_avg12[base_c]
            quote_flows = df_avg12[quote_c]
            spreads = base_flows - quote_flows
            ratings = spreads.apply(rate_fx_spread)

            pair_res = {}
            for ym in date_strs:
                if ym >= '2023-01': # 48 months from 2023-01 to 2026-12
                    b_f = base_flows.loc[ym] if ym in base_flows.index else None
                    q_f = quote_flows.loc[ym] if ym in quote_flows.index else None
                    sp = spreads.loc[ym] if ym in spreads.index else None
                    rt = ratings.loc[ym] if ym in ratings.index else None
                    pair_res[ym] = {
                        'base_flow': float(b_f) if pd.notna(b_f) else None,
                        'quote_flow': float(q_f) if pd.notna(q_f) else None,
                        'spread': float(sp) if pd.notna(sp) else None,
                        'rating': int(rt) if pd.notna(rt) else 0,
                    }
            pair_results[pair_name] = pair_res

    country_results = {}
    country_flow_results = {}
    for c_name in df_raw.columns:
        c_raw = {}
        c_flow = {}
        for ym in date_strs:
            v = df_raw.loc[ym, c_name]
            f = df_avg12.loc[ym, c_name]
            c_raw[ym] = float(v) if pd.notna(v) else None
            c_flow[ym] = float(f) if pd.notna(f) else None
        country_results[c_name] = c_raw
        country_flow_results[c_name] = c_flow

    return df_raw, df_avg12, pair_results, country_results, country_flow_results

if __name__ == "__main__":
    df_raw, df_avg12, pair_res, c_raw, c_flow = build_fx_dataset()
    print("\nSAMPLE FX DIFFERENTIAL & RATINGS FOR 2026-06:")
    for p in ['EUR/USD', 'USD/JPY', 'GBP/USD', 'AUD/USD', 'USD/CAD', 'USD/CHF']:
        pt = pair_res[p]['2026-06']
        print(f"{p:<9}: Base Flow={pt['base_flow']:+9.2f}M | Quote Flow={pt['quote_flow']:+9.2f}M | Spread={pt['spread']:+9.2f}M | Rating={pt['rating']:+3d}")

    print("\nSAMPLE FX DIFFERENTIAL & RATINGS FOR 2026-09:")
    for p in ['EUR/USD', 'USD/JPY', 'GBP/USD', 'AUD/USD', 'USD/CAD', 'USD/CHF']:
        pt = pair_res[p]['2026-09']
        print(f"{p:<9}: Base Flow={pt['base_flow']:+9.2f}M | Quote Flow={pt['quote_flow']:+9.2f}M | Spread={pt['spread']:+9.2f}M | Rating={pt['rating']:+3d}")
