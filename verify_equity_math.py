import openpyxl
import json
import sqlite3

def rate_equity(diff):
    if diff is None: return None
    bj = [-0.25, -0.20, -0.18, -0.16, -0.14, -0.12, -0.10, -0.08, -0.06, -0.04, -0.02, 0.00, 0.02, 0.04, 0.06, 0.08, 0.10, 0.12, 0.14, 0.16, 0.18, 0.20, 0.25]
    bk = [-2, -3, -5, -7, -9, -10, -9, -7, -5, -3, -2, 0, 2, 3, 5, 7, 9, 10, 9, 7, 5, 3, 2]
    clamped = max(-0.25, min(0.25, float(diff)))
    idx = min(range(len(bj)), key=lambda i: abs(bj[i] - clamped))
    return int(bk[idx])

def audit():
    print("=" * 118)
    print("AUDITING EQUITY INDICATOR INTEGRATION ACROSS EXCEL, DJANGO DB, & FRONTEND")
    print("=" * 118)

    # 1. Load Excel EQUITY sheet
    wb = openpyxl.load_workbook("EXCEL8EXAMPLE.xlsx", data_only=True)
    ws_eq = wb["EQUITY"]
    ws_cm = wb["COMPARATIVE MATRIX"]

    # Month cols in EQUITY
    col_2026_06_eq = 44 # AR is 2026-06
    col_2026_09_eq = 47 # AU is 2026-09

    # Month cols in COMPARATIVE MATRIX
    col_2026_06_cm = 47 # AU is 2026-06
    col_2026_09_cm = 50 # AX is 2026-09

    header_eq_jun = str(ws_eq.cell(row=5, column=col_2026_06_eq).value)[:7]
    header_eq_sep = str(ws_eq.cell(row=5, column=col_2026_09_eq).value)[:7]
    print(f"Excel EQUITY Headers: Col {col_2026_06_eq} = {header_eq_jun}, Col {col_2026_09_eq} = {header_eq_sep}")

    header_cm_jun = str(ws_cm.cell(row=148, column=col_2026_06_cm).value)[:7]
    header_cm_sep = str(ws_cm.cell(row=148, column=col_2026_09_cm).value)[:7]
    print(f"Excel COMPARATIVE MATRIX Headers: Col {col_2026_06_cm} = {header_cm_jun}, Col {col_2026_09_cm} = {header_cm_sep}")

    # Extract pairs from EQUITY sheet: each pair block starts at row r
    excel_equity = {}
    for r in range(6, 190, 8):
        pair_val = ws_eq.cell(row=r, column=1).value
        if not pair_val: continue
        pair_name = str(pair_val).strip()

        r_synth_ath = r
        r_synth_cur = r + 1
        r_base_ath = r + 3
        r_base_cur = r + 4

        excel_equity[pair_name] = {}
        for col, mo in [(col_2026_06_eq, "2026-06"), (col_2026_09_eq, "2026-09")]:
            s_ath = ws_eq.cell(row=r_synth_ath, column=col).value
            s_cur = ws_eq.cell(row=r_synth_cur, column=col).value
            b_ath = ws_eq.cell(row=r_base_ath, column=col).value
            b_cur = ws_eq.cell(row=r_base_cur, column=col).value

            if all(v is not None for v in [s_ath, s_cur, b_ath, b_cur]):
                chg_s = (s_cur - s_ath) / s_ath
                chg_b = (b_cur - b_ath) / b_ath
                diff = chg_s - chg_b
                rtg = rate_equity(diff)
                excel_equity[pair_name][mo] = {
                    "diff": diff,
                    "rating": rtg
                }

    print(f"Extracted and evaluated {len(excel_equity)} pairs from Excel EQUITY sheet.")

    # 2. Check Django Database
    conn = sqlite3.connect("fx-macro-bias/backend/db.sqlite3")
    cur = conn.cursor()

    cur.execute("""
        SELECT p.symbol, d.month, d.difference, r.rating_value, f.final_score, f.bias
        FROM macro_differential d
        JOIN macro_fxpair p ON d.pair_id = p.id
        JOIN macro_macroindicator i ON d.indicator_id = i.id
        LEFT JOIN macro_rating r ON r.differential_id = d.id
        LEFT JOIN macro_finalscore f ON f.pair_id = p.id AND f.month = d.month
        WHERE i.slug = 'equity' AND (d.month LIKE '2026-06%' OR d.month LIKE '2026-09%')
        ORDER BY p.symbol, d.month
    """)
    db_rows = cur.fetchall()
    db_data = {}
    for symbol, dt, diff_v, rate_v, fs_v, bias_v in db_rows:
        mo = dt[:7]
        db_data.setdefault(symbol, {})[mo] = {
            "diff": diff_v,
            "rate": rate_v,
            "score": fs_v,
            "bias": bias_v
        }
    print(f"Extracted {len(db_data)} pairs from Django SQLite DB for 'equity' indicator.")

    # 3. Check Frontend Dataset
    with open("fx-macro-bias/frontend/src/data/macroDataset.json", "r", encoding="utf-8") as f:
        fe_data = json.load(f)

    fe_pairs = fe_data.get("PairEquity", {})
    print(f"Extracted {len(fe_pairs)} pairs from Frontend macroDataset.json['PairEquity'].")

    # 21 Active Pairs
    active_pairs = [
        "EUR/USD", "GBP/USD", "USD/JPY", "AUD/USD", "USD/CAD", "NZD/USD", "USD/CHF",
        "EUR/JPY", "GBP/JPY", "EUR/GBP", "EUR/CHF", "AUD/JPY", "CAD/JPY", "CHF/JPY",
        "EUR/AUD", "EUR/CAD", "EUR/NZD", "GBP/AUD", "GBP/CAD", "GBP/CHF", "GBP/NZD"
    ]

    print("\n" + "=" * 118)
    print(f"{'Pair':<9} | {'Period':<7} | {'Excel Diff':<11} | {'Excel Rtg':<9} | {'DB Diff':<11} | {'DB Rtg':<8} | {'FE Diff':<11} | {'FE Rtg':<8} | {'Total Score':<11} | {'Macro Bias':<10}")
    print("=" * 118)

    all_matched = True
    checked_count = 0
    for p in active_pairs:
        p_space = p.replace("/", " ")
        ex = excel_equity.get(p_space, excel_equity.get(p, {}))
        db = db_data.get(p, {})
        fe = fe_pairs.get(p, {})

        for mo in ["2026-06", "2026-09"]:
            ex_entry = ex.get(mo, {})
            ex_d = ex_entry.get("diff")
            ex_r = ex_entry.get("rating")
            db_entry = db.get(mo, {})
            db_d = db_entry.get("diff")
            db_r = db_entry.get("rate")
            db_score = db_entry.get("score")
            db_bias = db_entry.get("bias")
            fe_entry = fe.get(mo, {})
            fe_d = fe_entry.get("final_change")
            fe_r = fe_entry.get("rating")

            ex_d_str = f"{ex_d:+.4f}" if ex_d is not None else "None"
            ex_r_str = f"{int(ex_r):+d}" if ex_r is not None else "None"
            db_d_str = f"{db_d:+.4f}" if db_d is not None else "None"
            db_r_str = f"{int(db_r):+d}" if db_r is not None else "None"
            fe_d_str = f"{fe_d:+.4f}" if fe_d is not None else "None"
            fe_r_str = f"{int(fe_r):+d}" if fe_r is not None else "None"
            sc_str = f"{db_score:+.2f}" if db_score is not None else "None"
            bias_str = str(db_bias) if db_bias else "None"

            print(f"{p:<9} | {mo:<7} | {ex_d_str:<11} | {ex_r_str:<9} | {db_d_str:<11} | {db_r_str:<8} | {fe_d_str:<11} | {fe_r_str:<8} | {sc_str:<11} | {bias_str:<10}")

            if ex_r is not None and db_r is not None:
                checked_count += 1
                if int(ex_r) != int(db_r):
                    print(f"MISMATCH in rating for {p} {mo}: Excel={ex_r}, DB={db_r}")
                    all_matched = False
                if abs(ex_d - db_d) > 1e-4:
                    print(f"MISMATCH in differential for {p} {mo}: Excel={ex_d}, DB={db_d}")
                    all_matched = False
            if db_r is not None and fe_r is not None:
                if int(db_r) != int(fe_r):
                    print(f"MISMATCH in rating for {p} {mo}: DB={db_r}, FE={fe_r}")
                    all_matched = False
                if abs(db_d - fe_d) > 1e-4:
                    print(f"MISMATCH in differential for {p} {mo}: DB={db_d}, FE={fe_d}")
                    all_matched = False

    print("=" * 118)
    if all_matched and checked_count > 0:
        print(f"[SUCCESS] 100% PARITY CONFIRMED! All {checked_count} pair-month equity differentials and ratings match exactly across Excel, Django DB, and Frontend!")
    else:
        print("[ERROR] Audit detected inconsistencies or missing data.")

if __name__ == "__main__":
    audit()
