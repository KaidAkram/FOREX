import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np
import datetime
import equity_engine

def update_excel_workbook():
    print("--- 1. Generating Equity Dataset via equity_engine ---")
    pair_results, country_results = equity_engine.build_equity_dataset(use_cache=True)

    print("--- 2. Updating EXCEL8EXAMPLE.xlsx ---")
    wb = openpyxl.load_workbook('EXCEL8EXAMPLE.xlsx', data_only=False)
    ws_eq = wb['EQUITY']
    ws_cm = wb['COMPARATIVE MATRIX']

    # 48 months from 2023-01 to 2026-12
    all_months = []
    for y in range(2023, 2027):
        for m in range(1, 13):
            all_months.append(f"{y}-{m:02d}")

    # Set dates in Row 5 of EQUITY sheet from Col 3 (C) to Col 50 (AX)
    for idx, ym in enumerate(all_months):
        col = idx + 3
        y, m = ym.split('-')
        ws_eq.cell(5, col).value = datetime.datetime(int(y), int(m), 1, 0, 0)

    # 23 pairs in EQUITY sheet
    pairs_order = [
        ('USD JPY', 6),
        ('USD CHF', 14),
        ('USD CAD', 22),
        ('EUR USD', 30),
        ('EUR JPY', 38),
        ('EUR CHF', 46),
        ('EUR GBP', 54),
        ('EUR CAD', 62),
        ('EUR AUD', 70),
        ('EUR NZD', 78),
        ('GBP USD', 86),
        ('GBP JPY', 94),
        ('GBP CHF', 102),
        ('GBP CAD', 110),
        ('GBP AUD', 118),
        ('GBP NZD', 126),
        ('CHF JPY', 134),
        ('AUD USD', 142),
        ('AUD CHF', 150),
        ('AUD JPY', 158),
        ('AUD NZD', 166),
        ('AUD CAD', 174),
        ('CAD JPY', 182),
    ]

    base_index_labels = {
        'USD': 'SPX',
        'EUR': 'STOXX50',
        'CAD': 'TSX',
        'JPY': 'NIKKEI',
        'GBP': 'TSE UK',
        'AUD': 'ASX200',
        'CHF': 'SMI',
        'NZD': 'NZ50',
    }

    rating_rows = {}

    for pair_name, r in pairs_order:
        p_data = pair_results[pair_name]
        base_code, _ = equity_engine.FX_TICKERS[pair_name]
        idx_lbl = base_index_labels[base_code]

        # Fix row labels
        ws_eq.cell(r, 1).value = pair_name
        ws_eq.cell(r, 2).value = 'All time High (CL)'
        ws_eq.cell(r+1, 2).value = 'current price'
        ws_eq.cell(r+2, 2).value = '% CHANGE'
        ws_eq.cell(r+3, 2).value = f'All time High {idx_lbl}'
        ws_eq.cell(r+4, 2).value = f'current price {idx_lbl}'
        ws_eq.cell(r+5, 2).value = 'original index % change'
        ws_eq.cell(r+6, 2).value = 'real % change'
        ws_eq.cell(r+7, 2).value = 'Rating'

        rating_rows[pair_name] = r + 7

        for idx, ym in enumerate(all_months):
            col = idx + 3
            col_ltr = get_column_letter(col)
            pt = p_data[ym]

            # Row r: Synth ATH
            ws_eq.cell(r, col).value = pt['synth_ath']
            # Row r+1: Synth Current Price
            ws_eq.cell(r+1, col).value = pt['synth_cur']
            # Row r+2: % CHANGE formula: =+(C7-C6)/C6
            ws_eq.cell(r+2, col).value = f"=+({col_ltr}{r+1}-{col_ltr}{r})/{col_ltr}{r}"
            # Row r+3: Base Index ATH
            ws_eq.cell(r+3, col).value = pt['base_ath']
            # Row r+4: Base Index Current Price
            ws_eq.cell(r+4, col).value = pt['base_cur']
            # Row r+5: original index % change formula: =+(C10-C9)/C9
            ws_eq.cell(r+5, col).value = f"=+({col_ltr}{r+4}-{col_ltr}{r+3})/{col_ltr}{r+3}"
            # Row r+6: real % change formula: =+C11-C8 (%a - %b: original index % change - synth % change)
            ws_eq.cell(r+6, col).value = f"=+{col_ltr}{r+5}-{col_ltr}{r+2}"
            # Row r+7: Rating formula
            ws_eq.cell(r+7, col).value = f'=_xlfn.XLOOKUP(MIN(ABS($BJ$7:$BJ$29 - MEDIAN(MIN($BJ$7:$BJ$29), MAX($BJ$7:$BJ$29), {col_ltr}{r+6}))), ABS($BJ$7:$BJ$29 - MEDIAN(MIN($BJ$7:$BJ$29), MAX($BJ$7:$BJ$29), {col_ltr}{r+6})), $BK$7:$BK$29, "Error", 0)'

    print("   -> Updated all 23 pairs across 48 months in EQUITY sheet!")

    # --- 3. Update COMPARATIVE MATRIX sheet ---
    print("--- 3. Updating COMPARATIVE MATRIX sheet ---")
    cm_pairs_order = [
        ('USD JPY', 149),
        ('USD CHF', 150),
        ('USD CAD', 151),
        ('EUR USD', 152),
        ('EUR JPY', 153),
        ('EUR CHF', 154),
        ('EUR GBP', 155),
        ('EUR CAD', 156),
        ('EUR AUD', 157),
        ('EUR NZD', 158),
        ('GBP USD', 159),
        ('GBP JPY', 160),
        ('GBP CHF', 161),
        ('GBP CAD', 162),
        ('GBP AUD', 163),
        ('GBP NZD', 164),
        ('CHF JPY', 165),
        ('AUD USD', 166),
        ('AUD CHF', 167),
        ('AUD JPY', 168),
        ('AUD NZD', 169),
        ('AUD CAD', 170),
        ('CAD JPY', 171),
    ]

    # For each pair and month, link CM cell to EQUITY rating cell
    for p_name, cm_row in cm_pairs_order:
        eq_rat_row = rating_rows[p_name]
        ws_cm.cell(cm_row, 5).value = p_name

        for idx, ym in enumerate(all_months):
            cm_col = idx + 6 # Col 6 is F (2023-01)
            eq_col = idx + 3 # Col 3 is C (2023-01)
            eq_col_ltr = get_column_letter(eq_col)

            # Set formula linking to EQUITY sheet rating
            ws_cm.cell(cm_row, cm_col).value = f"=EQUITY!{eq_col_ltr}{eq_rat_row}"

    print("   -> Linked all 23 pairs in COMPARATIVE MATRIX to EQUITY ratings!")

    wb.save('EXCEL8EXAMPLE.xlsx')
    print("[SUCCESS] EXCEL8EXAMPLE.xlsx updated and saved successfully!")

if __name__ == "__main__":
    update_excel_workbook()
