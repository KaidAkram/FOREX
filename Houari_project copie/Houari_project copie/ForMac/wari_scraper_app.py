import customtkinter as ctk
import threading
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
import pandas as pd
import time
import os
import datetime
import urllib.request
import json
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# --- Setup Modern UI Theme ---
ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

# The exact list of countries provided, formatted for display
VALID_COUNTRIES_DISPLAY = [
    "United States", "China", "Euro Area", "Germany", "Japan", "India", "United Kingdom",
    "France", "Italy", "Canada", "Brazil", "Russia", "South Korea", "Mexico", "Australia",
    "Spain", "Indonesia", "Turkey", "Saudi Arabia", "Netherlands", "Switzerland", "Poland",
    "Taiwan", "Belgium", "Argentina", "Sweden", "Ireland", "Singapore", "Israel",
    "United Arab Emirates", "Thailand", "Austria", "Norway", "Vietnam", "Philippines",
    "Bangladesh", "Iran", "Denmark", "Malaysia", "Colombia", "Hong Kong", "South Africa",
    "Egypt", "Romania", "Pakistan", "Czech Republic", "Chile", "Portugal", "Finland",
    "Peru", "Kazakhstan", "Iraq", "Algeria", "New Zealand", "Greece", "Hungary", "Qatar",
    "Ukraine", "Nigeria", "Kuwait", "Morocco", "Ethiopia", "Slovakia", "Puerto Rico",
    "Kenya", "Ecuador", "Dominican Republic", "Uzbekistan", "Guatemala", "Bulgaria",
    "Venezuela", "Cuba", "Oman", "Sri Lanka", "Costa Rica", "Croatia", "Luxembourg",
    "Serbia", "Ivory Coast", "Panama", "Lithuania", "Ghana", "Uruguay", "Angola",
    "Tanzania", "Belarus", "Azerbaijan", "Myanmar", "Slovenia", "Congo", "Turkmenistan",
    "Uganda", "Tunisia", "Jordan", "Cameroon", "Bolivia", "Macau", "Sudan", "Bahrain",
    "Libya", "Cambodia", "Latvia", "Paraguay", "Zimbabwe", "Nepal", "Estonia", "Honduras",
    "Cyprus", "El Salvador", "Georgia", "Iceland", "Papua New Guinea", "Senegal",
    "Bosnia and Herzegovina", "Albania", "Mali", "Armenia", "Zambia", "Trinidad and Tobago",
    "Guyana", "Haiti", "Guinea", "Malta", "Mongolia", "Burkina Faso", "Mozambique", "Chad",
    "Benin", "Gabon", "Nicaragua", "Niger", "Lebanon", "Jamaica", "Syria", "Yemen",
    "Botswana", "Moldova", "Macedonia", "Madagascar", "Kyrgyzstan", "Laos", "Afghanistan",
    "Bahamas", "Republic of the Congo", "Mauritius", "Brunei", "Rwanda", "Palestine",
    "Tajikistan", "Namibia", "Equatorial Guinea", "Somalia", "South Sudan", "Monaco",
    "Mauritania", "Malawi", "Kosovo", "Togo", "New Caledonia", "Montenegro", "Liechtenstein",
    "Sierra Leone", "Maldives", "Cayman Islands", "Fiji", "Liberia", "Suriname",
    "Swaziland", "Djibouti", "Gambia", "Central African Republic", "Cape Verde", "Bhutan",
    "Burundi", "Comoros", "Guinea Bissau", "East Timor", "Eritrea", "Lesotho", "Seychelles",
    "Sao Tome and Principe"
]

# Sort alphabetically for a better user experience
VALID_COUNTRIES_DISPLAY.sort()

# The 4 Core Macro Indicators tracked by the platform
INDICATORS = {
    "GDP Growth": "gdp-growth-annual",
    "Inflation": "inflation-cpi",
    "Interest Rate": "interest-rate",
    "FX Reserves": "imf"
}

# ISO/SDMX Mapping for IMF API queries
IMF_COUNTRY_MAP = {
    "albania": "ALB",
    "algeria": "DZA",
    "angola": "AGO",
    "argentina": "ARG",
    "armenia": "ARM",
    "australia": "AUS",
    "austria": "AUT",
    "azerbaijan": "AZE",
    "bahamas": "BHS",
    "bahrain": "BHR",
    "bangladesh": "BGD",
    "belarus": "BLR",
    "belgium": "BEL",
    "bhutan": "BTN",
    "brazil": "BRA",
    "brunei": "BRN",
    "bulgaria": "BGR",
    "cambodia": "KHM",
    "canada": "CAN",
    "cape-verde": "CPV",
    "chile": "CHL",
    "china": "CHN",
    "colombia": "COL",
    "comoros": "COM",
    "costa-rica": "CRI",
    "croatia": "HRV",
    "cyprus": "CYP",
    "czech-republic": "CZE",
    "denmark": "DNK",
    "djibouti": "DJI",
    "dominican-republic": "DOM",
    "east-timor": "TLS",
    "ecuador": "ECU",
    "egypt": "EGY",
    "el-salvador": "SLV",
    "estonia": "EST",
    "euro-area": "G163",
    "finland": "FIN",
    "france": "FRA",
    "georgia": "GEO",
    "germany": "DEU",
    "ghana": "GHA",
    "greece": "GRC",
    "guatemala": "GTM",
    "guyana": "GUY",
    "haiti": "HTI",
    "honduras": "HND",
    "hong-kong": "HKG",
    "hungary": "HUN",
    "iceland": "ISL",
    "india": "IND",
    "indonesia": "IDN",
    "iraq": "IRQ",
    "ireland": "IRL",
    "israel": "ISR",
    "italy": "ITA",
    "japan": "JPN",
    "jordan": "JOR",
    "kazakhstan": "KAZ",
    "kenya": "KEN",
    "kuwait": "KWT",
    "kyrgyzstan": "KGZ",
    "laos": "LAO",
    "latvia": "LVA",
    "libya": "LBY",
    "lithuania": "LTU",
    "luxembourg": "LUX",
    "macau": "MAC",
    "macedonia": "MKD",
    "malaysia": "MYS",
    "maldives": "MDV",
    "malta": "MLT",
    "mauritius": "MUS",
    "mexico": "MEX",
    "moldova": "MDA",
    "mongolia": "MNG",
    "morocco": "MAR",
    "namibia": "NAM",
    "netherlands": "NLD",
    "new-zealand": "NZL",
    "nicaragua": "NIC",
    "nigeria": "NGA",
    "norway": "NOR",
    "pakistan": "PAK",
    "panama": "PAN",
    "peru": "PER",
    "philippines": "PHL",
    "poland": "POL",
    "portugal": "PRT",
    "qatar": "QAT",
    "romania": "ROU",
    "rwanda": "RWA",
    "saudi-arabia": "SAU",
    "serbia": "SRB",
    "singapore": "SGP",
    "slovakia": "SVK",
    "slovenia": "SVN",
    "south-africa": "ZAF",
    "south-korea": "KOR",
    "spain": "ESP",
    "suriname": "SUR",
    "sweden": "SWE",
    "switzerland": "CHE",
    "taiwan": "TWN",
    "thailand": "THA",
    "trinidad-and-tobago": "TTO",
    "turkey": "TUR",
    "ukraine": "UKR",
    "united-arab-emirates": "ARE",
    "united-kingdom": "GBR",
    "united-states": "USA",
    "uruguay": "URY",
    "uzbekistan": "UZB",
    "vietnam": "VNM",
    "zimbabwe": "ZWE"
}

# The Core FX Pairs for Differential & Macro Rating Calculation
PAIRS_DEF = [
    ("EUR/USD", "Euro Area", "United States"),
    ("GBP/USD", "United Kingdom", "United States"),
    ("USD/JPY", "United States", "Japan"),
    ("AUD/USD", "Australia", "United States"),
    ("USD/CAD", "United States", "Canada"),
    ("NZD/USD", "New Zealand", "United States"),
    ("USD/CHF", "United States", "Switzerland"),
    ("EUR/JPY", "Euro Area", "Japan"),
    ("GBP/JPY", "United Kingdom", "Japan"),
    ("EUR/GBP", "Euro Area", "United Kingdom"),
    ("EUR/CHF", "Euro Area", "Switzerland"),
    ("AUD/JPY", "Australia", "Japan")
]

def rate_interest_rate(diff):
    import numpy as np
    if pd.isna(diff): return np.nan
    # Exact lookup table from EXCEL8EXAMPLE.xlsx Interest Rates data sheet (Cols AZ & BA)
    az = [-7.0, -6.5, -6.0, -5.5, -5.0, -4.5, -4.0, -3.5, -3.0, -2.5, -2.0, -1.5, -1.0, -0.5, 0.0, 0.5, 1.0, 1.5, 2.0, 2.5, 3.0, 3.5, 4.0, 4.5, 5.0, 5.5, 6.0, 6.5, 7.0]
    ba = [-3,   -4,   -5,   -6,   -10,  -9,   -8,   -7,   -6,   -5,   -4,   -3,   -2,   -1,   0,   1,   2,   3,   4,   5,   6,   7,   8,   9,   10,   6,   5,   4,   3]
    clamped = max(-7.0, min(7.0, float(diff)))
    idx = min(range(len(az)), key=lambda i: abs(az[i] - clamped))
    return int(ba[idx])

def rate_cpi(diff):
    import numpy as np
    if pd.isna(diff): return np.nan
    # Exact lookup table from EXCEL8EXAMPLE.xlsx CPI data sheet (Cols AY & AZ)
    ay = [-2.0, -1.75, -1.5, -1.25, -1.0, -0.75, -0.5, -0.25, 0.0, 0.25, 0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0]
    az = [2,    3,     4,    6,     10,   8,     6,    4,     0,   -4,   -6,  -8,   -10, 6,    8,   9,    10]
    clamped = max(-2.0, min(2.0, float(diff)))
    idx = min(range(len(ay)), key=lambda i: abs(ay[i] - clamped))
    return int(az[idx])

def rate_gdp(diff):
    import numpy as np
    if pd.isna(diff): return np.nan
    # Exact formula from EXCEL8EXAMPLE.xlsx GDP data sheet: 2 * diff clamped to [-10, +10]
    clamped = max(-5.0, min(5.0, float(diff)))
    return int(round(clamped * 2))

def rate_fx_spread(spread):
    import numpy as np
    if pd.isna(spread): return np.nan
    # Exact lookup table from EXCEL8EXAMPLE.xlsx FX RESERVE DATA sheet (Cols BK & BL)
    bk = [-2000 + 200*i for i in range(21)]
    bl = [10 - i for i in range(21)]
    clamped = max(-2000.0, min(2000.0, float(spread)))
    idx = min(range(len(bk)), key=lambda i: abs(bk[i] - clamped))
    return int(bl[idx])

def compute_post_scraping_math(matrices_val, matrices_color, log_fn=print):
    """
    Computes mathematical models following EXCEL8EXAMPLE specification:
    1. Differentials & Rating for Interest Rate, Inflation (CPI), and GDP Growth
    2. 12-Month rolling average monthly flow and spreads for FX Reserves
    3. Final COMPARATIVE MATRIX sheet compiling composite scores and macro bias
    """
    new_val = dict(matrices_val)
    new_color = dict(matrices_color)
    pair_ratings = {}

    # 1. Process Interest Rate Differentials
    if "Interest Rate" in matrices_val:
        df_ir = matrices_val["Interest Rate"].copy().set_index("Country")
        timeframes = [c for c in df_ir.columns if c != "Country"]
        diff_rows = []
        rating_rows = []
        for pair_name, base_c, quote_c in PAIRS_DEF:
            if base_c in df_ir.index and quote_c in df_ir.index:
                b_series = pd.to_numeric(df_ir.loc[base_c, timeframes], errors='coerce')
                q_series = pd.to_numeric(df_ir.loc[quote_c, timeframes], errors='coerce')
                diff_series = (b_series - q_series).round(2)
                rate_series = diff_series.apply(rate_interest_rate)
                diff_rows.append({"Pair": pair_name, **diff_series.to_dict()})
                rating_rows.append({"Pair": pair_name, **rate_series.to_dict()})
                pair_ratings.setdefault(pair_name, {})["Interest Rate"] = rate_series

        if diff_rows:
            new_val["Diff - Interest Rate"] = pd.DataFrame(diff_rows)
            new_val["Rating - Interest Rate"] = pd.DataFrame(rating_rows)
            log_fn("  -> ✅ Calculated Interest Rate Differentials & Ratings")

    # 2. Process Inflation Differentials
    cpi_key = "Inflation" if "Inflation" in matrices_val else ("CPI Inflation" if "CPI Inflation" in matrices_val else None)
    if cpi_key:
        df_cpi = matrices_val[cpi_key].copy().set_index("Country")
        timeframes = [c for c in df_cpi.columns if c != "Country"]
        diff_rows = []
        rating_rows = []
        for pair_name, base_c, quote_c in PAIRS_DEF:
            if base_c in df_cpi.index and quote_c in df_cpi.index:
                b_series = pd.to_numeric(df_cpi.loc[base_c, timeframes], errors='coerce')
                q_series = pd.to_numeric(df_cpi.loc[quote_c, timeframes], errors='coerce')
                diff_series = (b_series - q_series).round(2)
                rate_series = diff_series.apply(rate_cpi)
                diff_rows.append({"Pair": pair_name, **diff_series.to_dict()})
                rating_rows.append({"Pair": pair_name, **rate_series.to_dict()})
                pair_ratings.setdefault(pair_name, {})["Inflation"] = rate_series

        if diff_rows:
            new_val["Diff - Inflation"] = pd.DataFrame(diff_rows)
            new_val["Rating - Inflation"] = pd.DataFrame(rating_rows)
            log_fn("  -> ✅ Calculated Inflation Differentials & Ratings")

    # 3. Process GDP Growth Differentials
    if "GDP Growth" in matrices_val:
        df_gdp = matrices_val["GDP Growth"].copy().set_index("Country")
        timeframes = [c for c in df_gdp.columns if c != "Country"]
        diff_rows = []
        rating_rows = []
        for pair_name, base_c, quote_c in PAIRS_DEF:
            if base_c in df_gdp.index and quote_c in df_gdp.index:
                b_series = pd.to_numeric(df_gdp.loc[base_c, timeframes], errors='coerce')
                q_series = pd.to_numeric(df_gdp.loc[quote_c, timeframes], errors='coerce')
                diff_series = (b_series - q_series).round(2)
                rate_series = diff_series.apply(rate_gdp)
                diff_rows.append({"Pair": pair_name, **diff_series.to_dict()})
                rating_rows.append({"Pair": pair_name, **rate_series.to_dict()})
                pair_ratings.setdefault(pair_name, {})["GDP Growth"] = rate_series

        if diff_rows:
            new_val["Diff - GDP Growth"] = pd.DataFrame(diff_rows)
            new_val["Rating - GDP Growth"] = pd.DataFrame(rating_rows)
            log_fn("  -> ✅ Calculated GDP Growth Differentials & Ratings")

    # 4. Process FX Reserves: 12-Month Rolling Flow & Spread
    if "FX Reserves" in matrices_val:
        df_fx = matrices_val["FX Reserves"].copy().set_index("Country")
        timeframes = [c for c in df_fx.columns if c != "Country"]
        df_fx_num = df_fx[timeframes].apply(pd.to_numeric, errors='coerce')
        delta = df_fx_num.diff(axis=1)
        rolling_12m = delta.T.rolling(window=12, min_periods=1).mean().T.round(2)
        
        spread_rows = []
        rating_rows = []
        for pair_name, base_c, quote_c in PAIRS_DEF:
            if base_c in rolling_12m.index and quote_c in rolling_12m.index:
                b_series = rolling_12m.loc[base_c]
                q_series = rolling_12m.loc[quote_c]
                spread_series = (b_series - q_series).round(2)
                rate_series = spread_series.apply(rate_fx_spread)
                spread_rows.append({"Pair": pair_name, **spread_series.to_dict()})
                rating_rows.append({"Pair": pair_name, **rate_series.to_dict()})
                pair_ratings.setdefault(pair_name, {})["FX Reserves"] = rate_series

        if spread_rows:
            new_val["12M Avg Flow - FX Reserves"] = rolling_12m.reset_index()
            new_val["Spread - FX Reserves"] = pd.DataFrame(spread_rows)
            new_val["Rating - FX Reserves"] = pd.DataFrame(rating_rows)
            log_fn("  -> ✅ Calculated FX Reserves 12M Rolling Flow, Spreads & Ratings")

    # 5. Build COMPARATIVE MATRIX
    if pair_ratings:
        cm_rows = []
        for pair_name, ind_dict in pair_ratings.items():
            sample_series = next(iter(ind_dict.values()))
            t_cols = list(sample_series.index)
            pair_df = pd.DataFrame(ind_dict)
            total_score = pair_df.sum(axis=1, min_count=1)
            num_indicators = len(ind_dict)
            max_score = num_indicators * 10.0
            bias_pct = ((total_score / max_score) * 100).round(1)
            
            def get_bias(p):
                if pd.isna(p): return "N/A"
                if p >= 20.0: return "BULLISH"
                if p <= -20.0: return "BEARISH"
                return "NEUTRAL"
            
            regimes = bias_pct.apply(get_bias)

            for ind_name, r_series in ind_dict.items():
                cm_rows.append({"Pair": pair_name, "Metric": f"{ind_name} Rating", **r_series.to_dict()})
            cm_rows.append({"Pair": pair_name, "Metric": "TOTAL SCORE", **total_score.to_dict()})
            cm_rows.append({"Pair": pair_name, "Metric": "BIAS PERCENTAGE (%)", **bias_pct.to_dict()})
            cm_rows.append({"Pair": pair_name, "Metric": "MACRO REGIME", **regimes.to_dict()})
            cm_rows.append({"Pair": pair_name, "Metric": "---", **{c: "" for c in t_cols}})

        new_val["COMPARATIVE MATRIX"] = pd.DataFrame(cm_rows)
        log_fn("  -> 🏆 Assembled Final COMPARATIVE MATRIX & Macro Bias Regimes")

    return new_val, new_color

class WariScraperApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("WARI Extraction Platform")
        self.geometry("750x740")
        self.is_extracting = False
        
        # --- UI Layout ---
        self.title_label = ctk.CTkLabel(self, text="WARI Scraper", font=ctk.CTkFont(size=24, weight="bold"))
        self.title_label.pack(pady=(20, 5), padx=20, anchor="w")
        
        self.subtitle_label = ctk.CTkLabel(self, text="Trading Economics & IMF SDMX Forecast Extractor", text_color="gray", font=ctk.CTkFont(size=14))
        self.subtitle_label.pack(pady=(0, 10), padx=20, anchor="w")

        # --- Search & Select Frame ---
        self.input_frame = ctk.CTkFrame(self)
        self.input_frame.pack(pady=5, padx=20, fill="x")

        self.target_label = ctk.CTkLabel(self.input_frame, text="Select Target Countries:", font=ctk.CTkFont(weight="bold"))
        self.target_label.pack(pady=(10, 5), padx=15, anchor="w")

        # Search Bar
        self.search_entry = ctk.CTkEntry(self.input_frame, placeholder_text="🔍 Search for a country...", height=35)
        self.search_entry.pack(pady=(0, 5), padx=15, fill="x")
        self.search_entry.bind("<KeyRelease>", self.filter_countries)

        # Scrollable Checkbox Area
        self.scroll_frame = ctk.CTkScrollableFrame(self.input_frame, height=140)
        self.scroll_frame.pack(pady=(0, 10), padx=15, fill="x")

        self.checkbox_vars = {}
        self.checkbox_widgets = {}

        # Populate the scrollable frame
        for country in VALID_COUNTRIES_DISPLAY:
            var = ctk.BooleanVar(value=False)
            chk = ctk.CTkCheckBox(self.scroll_frame, text=country, variable=var)
            chk.pack(anchor="w", pady=3, padx=5)
            self.checkbox_vars[country] = var
            self.checkbox_widgets[country] = chk

            # Pre-select G10 benchmark countries as default
            if country in ["United States", "Euro Area", "Japan", "United Kingdom", "Australia", "Canada", "Switzerland", "New Zealand"]:
                var.set(True)

        # --- Options Frame ---
        self.options_frame = ctk.CTkFrame(self.input_frame, fg_color="transparent")
        self.options_frame.pack(pady=(5, 10), padx=15, fill="x")

        self.headless_var = ctk.BooleanVar(value=True)
        self.headless_checkbox = ctk.CTkCheckBox(self.options_frame, text="Stealth Mode (Hidden Browser)", variable=self.headless_var)
        self.headless_checkbox.pack(side="left", padx=(0, 20))

        self.autorun_var = ctk.BooleanVar(value=False)
        self.autorun_switch = ctk.CTkSwitch(self.options_frame, text="Enable Auto-Run (Every 20 Mins)", variable=self.autorun_var, command=self.on_autorun_toggle)
        self.autorun_switch.pack(side="left")

        # --- Log TextBox ---
        self.log_label = ctk.CTkLabel(self, text="Activity Log:", font=ctk.CTkFont(weight="bold"))
        self.log_label.pack(pady=(10, 0), padx=20, anchor="w")

        self.log_box = ctk.CTkTextbox(self, height=150, state="disabled", font=ctk.CTkFont(family="Consolas", size=12))
        self.log_box.pack(pady=(5, 10), padx=20, fill="both", expand=True)

        # --- Action Buttons ---
        self.btn_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.btn_frame.pack(pady=(0, 15), padx=20, fill="x")

        self.start_btn = ctk.CTkButton(self.btn_frame, text="Run Extraction", height=45, font=ctk.CTkFont(weight="bold"), command=self.start_scraping_thread)
        self.start_btn.pack(side="left", fill="x", expand=True, padx=(0, 10))

        self.clear_btn = ctk.CTkButton(self.btn_frame, text="Clear Logs", height=45, fg_color="#444444", hover_color="#333333", command=self.clear_logs)
        self.clear_btn.pack(side="right", fill="x", expand=True, padx=(10, 0))

    def filter_countries(self, event=None):
        search_query = self.search_entry.get().lower()
        for country, chk in self.checkbox_widgets.items():
            if search_query in country.lower():
                chk.pack(anchor="w", pady=3, padx=5)
            else:
                chk.pack_forget()

    def log(self, message):
        timestamp = datetime.datetime.now().strftime("%H:%M:%S")
        formatted_msg = f"[{timestamp}] {message}\n"
        
        def update_ui():
            self.log_box.configure(state="normal")
            self.log_box.insert("end", formatted_msg)
            self.log_box.see("end")
            self.log_box.configure(state="disabled")
            
        self.after(0, update_ui)

    def clear_logs(self):
        self.log_box.configure(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.configure(state="disabled")

    def on_autorun_toggle(self):
        if self.autorun_var.get():
            self.start_btn.configure(text="Start Auto-Run Sequence")
            self.log("🔁 Auto-Run activated. The scraper will repeat every 20 minutes.")
        else:
            if not self.is_extracting:
                self.start_btn.configure(text="Run Extraction")
            self.log("🛑 Auto-Run disabled.")

    def start_scraping_thread(self):
        if self.is_extracting:
            self.log("⚠️ Scraper is currently busy.")
            return

        selected_display_names = []
        for country, var in self.checkbox_vars.items():
            if var.get():
                selected_display_names.append(country)

        if not selected_display_names:
            self.log("⚠️ Please select at least one country from the list.")
            return
            
        countries = [c.lower().replace(" ", "-") for c in selected_display_names]

        self.is_extracting = True
        self.start_btn.configure(state="disabled", text="Extracting...")
        
        if self.autorun_var.get():
            thread = threading.Thread(target=self.autorun_loop, args=(countries, selected_display_names), daemon=True)
        else:
            thread = threading.Thread(target=self.single_run, args=(countries, selected_display_names), daemon=True)
            
        thread.start()

    def autorun_loop(self, countries, display_names):
        while self.autorun_var.get():
            self.after(0, lambda: self.start_btn.configure(text="Extracting..."))
            
            self.run_extraction(countries, display_names)
            
            if not self.autorun_var.get():
                break
                
            self.log("⏳ Extraction complete. Waiting for the next cycle...")
            
            total_seconds = 1200 
            for remaining in range(total_seconds, 0, -1):
                if not self.autorun_var.get():
                    self.log("🛑 Auto-Run loop stopped by user.")
                    break
                
                mins, secs = divmod(remaining, 60)
                timer_format = f"{mins:02d}:{secs:02d}"
                self.after(0, lambda t=timer_format: self.start_btn.configure(text=f"Next run in {t}..."))
                time.sleep(1)
                
        self.is_extracting = False
        self.after(0, lambda: self.start_btn.configure(state="normal", text="Run Extraction"))

    def single_run(self, countries, display_names):
        self.run_extraction(countries, display_names)
        self.is_extracting = False
        self.after(0, lambda: self.start_btn.configure(state="normal", text="Run Extraction"))

    def run_extraction(self, countries, display_names):
        self.log(f"🚀 Initializing extraction for: {', '.join(display_names)}")

        output_dir = os.path.join(os.getcwd(), "outputs")
        os.makedirs(output_dir, exist_ok=True)

        date_str = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M")
        name_prefix = "_".join(countries[:3])
        if len(countries) > 3: name_prefix += "_etc"
        filename = f"{name_prefix}_Forecast_{date_str}.xlsx"
        output_filepath = os.path.join(output_dir, filename)

        options = Options()
        if self.headless_var.get(): options.add_argument('--headless')
        options.add_argument('--disable-gpu')
        options.add_argument('--log-level=3') 
        options.add_argument('--window-size=1920,1080')
        options.add_argument('user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36')

        matrices_val = {}
        matrices_color = {}

        try:
            driver = webdriver.Chrome(options=options)
            
            # 1. Scrape TradingEconomics indicators
            for sheet_name, indicator in INDICATORS.items():
                if indicator == "imf":
                    continue

                self.log(f"📈 Starting Data Pull: {sheet_name} ({indicator})")
                all_extracted_data = []

                for country in countries:
                    url = f"https://tradingeconomics.com/{country}/{indicator}"
                    formatted_country_name = country.replace("-", " ").title()
                    self.log(f"  -> 🌍 Loading {formatted_country_name}...")
                    
                    try:
                        driver.get(url)
                        time.sleep(4)
                        
                        try:
                            driver.execute_script("var el = document.querySelector('a[href=\"#forecast\"]'); if(el) el.click();")
                            time.sleep(3)
                        except Exception:
                            pass
                        
                        # Corrected JavaScript Highcharts extraction:
                        # 1. Takes the active forecast chart (last chart instantiated)
                        # 2. Rejects null/undefined separator points
                        # 3. Normalizes timestamp to UTC Month (YYYY-MM-01)
                        # 4. Accurately detects grey forecast points
                        js_code = """
                        try {
                            var charts = Highcharts.charts.filter(c => !!c && c.series && c.series.length > 0);
                            if (charts.length === 0) return { error: "No chart found" };

                            var chart = charts[charts.length - 1];
                            var pts = [];

                            for (var s = 0; s < chart.series.length; s++) {
                                var series = chart.series[s];
                                if (!series.data) continue;

                                for (var i = 0; i < series.data.length; i++) {
                                    var p = series.data[i];
                                    if (p && p.y !== null && p.y !== undefined) {
                                        var xVal = p.category || p.name || p.x;
                                        var dateStr = '';
                                        if (!isNaN(xVal) && parseInt(xVal) > 1000000000000) {
                                            var d = new Date(parseInt(xVal));
                                            var yr = d.getUTCFullYear();
                                            var mo = String(d.getUTCMonth() + 1).padStart(2, '0');
                                            dateStr = yr + '-' + mo + '-01';
                                        } else {
                                            dateStr = String(xVal);
                                        }

                                        var ptColor = String(p.color || series.color || '').toLowerCase();
                                        if (p.graphic && p.graphic.element) {
                                            ptColor = String(p.graphic.element.getAttribute('fill') || ptColor).toLowerCase();
                                        }

                                        var isForecast = (s > 0 || ptColor.includes('205') || ptColor.includes('cdcdcd') || ptColor.includes('ccc') || ptColor.includes('d9d9d9'));
                                        pts.push({ "Timeframe": dateStr, "Value": p.y, "Is_Forecast": isForecast });
                                    }
                                }
                            }
                            return { points: pts };
                        } catch(err) {
                            return { error: err.message };
                        }
                        """
                        
                        res = driver.execute_script(js_code)
                        if isinstance(res, dict) and "points" in res:
                            pts = res["points"]
                            for pt in pts:
                                pt["Country"] = formatted_country_name
                                all_extracted_data.append(pt)
                            self.log(f"     ✅ Extracted {len(pts)} points for {formatted_country_name}")
                        else:
                            self.log(f"     ⚠️ Could not extract data for {formatted_country_name}. Skipping.")

                    except Exception as e:
                        self.log(f"     ❌ Error loading {formatted_country_name}: {e}")

                if all_extracted_data:
                    df = pd.DataFrame(all_extracted_data)
                    df_clean = df.drop_duplicates(subset=['Country', 'Timeframe'], keep='last')
                    timeframes = sorted(df_clean['Timeframe'].unique())
                    target_countries = [c.replace("-", " ").title() for c in countries]
                    
                    pivot_val = df_clean.pivot_table(index='Country', columns='Timeframe', values='Value', aggfunc='last')
                    pivot_val = pivot_val.reindex(index=target_countries, columns=timeframes)
                    
                    # For Interest Rates: forward-fill prevailing policy rate across months between meetings
                    if indicator == "interest-rate":
                        pivot_val = pivot_val.ffill(axis=1)
                        
                    pivot_color = df_clean.pivot_table(index='Country', columns='Timeframe', values='Is_Forecast', aggfunc='last')
                    pivot_color = pivot_color.reindex(index=target_countries, columns=timeframes).fillna(False)
                    
                    matrices_val[sheet_name] = pivot_val.reset_index()
                    matrices_color[sheet_name] = pivot_color.reset_index()

            # 2. Fetch FX Reserves from official IMF SDMX 3.0 API
            if "FX Reserves" in INDICATORS:
                self.log("🏦 Fetching Foreign Exchange Reserves from IMF SDMX 3.0 API...")
                imf_codes = []
                code_to_display = {}
                for country in countries:
                    code = IMF_COUNTRY_MAP.get(country)
                    if code:
                        imf_codes.append(code)
                        formatted_country = country.replace("-", " ").title()
                        code_to_display[code] = formatted_country
                        
                if imf_codes:
                    try:
                        imf_query_codes = "+".join(imf_codes)
                        base_url = f"https://api.imf.org/external/sdmx/3.0/data/dataflow/IMF.STA/IL/13.0.1/{imf_query_codes}.RXF11FX_REVS.USD.M"
                        # Properly URL-encode c[TIME_PERIOD] as c%5BTIME_PERIOD%5D to prevent Tomcat 400 Bad Request
                        query = "c%5BTIME_PERIOD%5D=ge:2018-01-01&attributes=all&detail=full&includeHistory=true"
                        url = f"{base_url}?{query}"
                        
                        req = urllib.request.Request(url, headers={
                            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
                            "Accept": "application/vnd.sdmx.data+json;version=1.0.0-wd, application/json"
                        })
                        with urllib.request.urlopen(req, timeout=25) as resp:
                            imf_raw = json.loads(resp.read().decode('utf-8'))
                            
                        s0 = imf_raw['data']['structures'][0]
                        countries_dim = s0['dimensions']['series'][0]['values']
                        resp_country_codes = [c['id'] for c in countries_dim]
                        time_periods = [t['value'] for t in s0['dimensions']['observation'][0]['values']]
                        series_dict = imf_raw['data']['dataSets'][0].get('series', {})
                        
                        fx_records = []
                        for s_key, s_data in series_dict.items():
                            c_idx = int(s_key.split(':')[0])
                            code = resp_country_codes[c_idx]
                            country_name = code_to_display.get(code, code)
                            obs = s_data.get('observations', {})
                            for t_idx_str, val_array in obs.items():
                                t_idx = int(t_idx_str)
                                period_str = time_periods[t_idx] # "2022-M01"
                                parts = period_str.split('-M')
                                formatted_date = f"{parts[0]}-{parts[1]}-01" if len(parts) == 2 else period_str
                                raw_val = val_array[0]
                                if raw_val is not None:
                                    val_m_usd = round(float(raw_val) / 1e6, 2)
                                    fx_records.append({
                                        "Country": country_name,
                                        "Timeframe": formatted_date,
                                        "Value": val_m_usd,
                                        "Is_Forecast": False
                                    })
                                    
                        if fx_records:
                            fx_df = pd.DataFrame(fx_records)
                            fx_clean = fx_df.drop_duplicates(subset=['Country', 'Timeframe'], keep='last')
                            timeframes = sorted(fx_clean['Timeframe'].unique())
                            target_countries = [code_to_display[c] for c in imf_codes if c in code_to_display]
                            
                            pivot_val = fx_clean.pivot_table(index='Country', columns='Timeframe', values='Value', aggfunc='last').reindex(index=target_countries, columns=timeframes).reset_index()
                            pivot_color = fx_clean.pivot_table(index='Country', columns='Timeframe', values='Is_Forecast', aggfunc='last').reindex(index=target_countries, columns=timeframes).fillna(False).reset_index()
                            
                            matrices_val["FX Reserves"] = pivot_val
                            matrices_color["FX Reserves"] = pivot_color
                            self.log(f"  -> ✅ Retrieved FX Reserves for {len(imf_codes)} countries ({len(timeframes)} months)")
                    except Exception as e:
                        self.log(f"  -> ⚠️ IMF API Fetch error: {e}")
                else:
                    self.log("  -> ℹ️ No IMF country codes mapped for selected countries.")

        except Exception as e:
            self.log(f"❌ Webdriver Error: {e}")
        finally:
            if 'driver' in locals(): driver.quit()

        # 3. Post-Scraping Math & Macro Comparative Transformation
        if matrices_val:
            self.log("🧮 Executing Post-Scraping Mathematical Pipeline & Comparative Matrix...")
            try:
                matrices_val, matrices_color = compute_post_scraping_math(matrices_val, matrices_color, self.log)
            except Exception as e:
                self.log(f"⚠️ Math calculation note: {e}")

        # 4. Export to multi-sheet Excel workbook with formatted styling
        if matrices_val:
            self.log("📊 Compiling Excel File with Multiple Sheets...")
            
            with pd.ExcelWriter(output_filepath, engine='openpyxl') as writer:
                for sheet_name, pivot_val in matrices_val.items():
                    pivot_val.to_excel(writer, index=False, sheet_name=sheet_name)
            
            self.log("🎨 Applying grey highlights to forecast cells...")
            
            wb = load_workbook(output_filepath)
            grey_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            
            for sheet_name, pivot_color in matrices_color.items():
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    for r_idx, row in enumerate(pivot_color.values):
                        for c_idx, is_forecast in enumerate(row):
                            if is_forecast == True: 
                                ws.cell(row=r_idx + 2, column=c_idx + 2).fill = grey_fill
                            
            wb.save(output_filepath)
            self.log(f"🎉 DONE! Saved to: outputs/{filename}")
        else:
            self.log("⚠️ No data was found for any indicator.")

if __name__ == "__main__":
    app = WariScraperApp()
    app.mainloop()